"""
Tests de la lecture à la demande d'un document joint et de la passerelle vers
l'espace d'exécution de code (27/10/2026).

Les fichiers sont de VRAIS documents fabriqués ici (PDF avec PyMuPDF, Word, Excel,
PowerPoint), lus par les vraies bibliothèques. Seuls les services distants sont
simulés : le stockage (retrouver_fichier) et l'OCR de secours. Ces tests vérifient
la logique, pas l'accord avec le stockage réel : à confirmer en production avec un
vrai fichier.
"""

import io
import sys
from types import SimpleNamespace
from unittest import mock

import pytest

sys.path.append("core")  # comme les autres tests

import core.lecture_document_conversation as lecteur  # noqa: E402
from core import extraction_documents  # noqa: E402


def _pdf(textes_par_page):
    """PDF réel : une page par élément ; une chaîne vide donne une page sans texte natif."""
    import fitz

    document = fitz.open()
    for texte in textes_par_page:
        page = document.new_page()
        if texte:
            page.insert_text((72, 72), texte)
    octets = document.tobytes()
    document.close()
    return octets


def _docx(paragraphes):
    import docx

    document = docx.Document()
    for p in paragraphes:
        document.add_paragraph(p)
    flux = io.BytesIO()
    document.save(flux)
    return flux.getvalue()


def _xlsx():
    import openpyxl

    classeur = openpyxl.Workbook()
    classeur.active.title = "Notes"
    classeur.active.append(["nom", "note"])
    classeur.active.append(["Awa", 17])
    classeur.create_sheet("Moyennes").append(["total", 15])
    flux = io.BytesIO()
    classeur.save(flux)
    return flux.getvalue()


def _pptx(titres):
    from pptx import Presentation

    presentation = Presentation()
    for titre in titres:
        diapositive = presentation.slides.add_slide(presentation.slide_layouts[5])
        diapositive.shapes.title.text = titre
    flux = io.BytesIO()
    presentation.save(flux)
    return flux.getvalue()


@pytest.fixture
def stockage(monkeypatch):
    """Simule le stockage : retrouver_fichier renvoie le fichier déposé par le test."""
    etat = {"ligne": None, "octets": None, "appels": []}

    def deposer(nom, octets, type_mime=""):
        etat["ligne"] = {"nom_fichier": nom, "type_mime": type_mime, "taille_octets": len(octets)}
        etat["octets"] = octets

    def faux_retrouver(url, user_id):
        etat["appels"].append((url, user_id))
        if etat["ligne"] is None or url != "https://stockage/fichier":
            return None
        return etat["ligne"], etat["octets"]

    monkeypatch.setattr(lecteur, "retrouver_fichier", faux_retrouver)
    etat["deposer"] = deposer
    return etat


def lire(**arguments):
    arguments.setdefault("url_fichier", "https://stockage/fichier")
    return lecteur.lire_document(user_id="u1", **arguments)


class TestInformations:
    def test_pdf_donne_le_nombre_de_pages(self, stockage):
        stockage["deposer"]("cours.pdf", _pdf(["a", "b", "c"]), "application/pdf")
        resultat = lire(mode="infos")
        assert "cours.pdf" in resultat and "3 pages" in resultat and "PDF" in resultat

    def test_pptx_donne_les_diapositives(self, stockage):
        stockage["deposer"]("expose.pptx", _pptx(["Un", "Deux"]))
        assert "2 diapositives" in lire(mode="infos")

    def test_xlsx_donne_les_feuilles(self, stockage):
        stockage["deposer"]("notes.xlsx", _xlsx())
        resultat = lire(mode="infos")
        assert "2 feuilles" in resultat and "Notes" in resultat and "Moyennes" in resultat

    def test_infos_ne_lit_aucun_texte(self, stockage, monkeypatch):
        stockage["deposer"]("cours.pdf", _pdf(["secret"]), "application/pdf")
        appel = mock.Mock()
        monkeypatch.setattr(lecteur, "lire_fichier", appel)
        assert "secret" not in lire(mode="infos")
        appel.assert_not_called()

    def test_fichier_illisible_ne_leve_pas(self):
        infos = lecteur.infos_document("faux.pdf", b"ceci n'est pas un pdf", "application/pdf")
        assert infos["famille"] == "pdf" and infos["pages"] is None


class TestPdfParPages:
    def test_intervalle_demande(self, stockage):
        stockage["deposer"]("c.pdf", _pdf([f"contenu{i}" for i in range(1, 11)]), "application/pdf")
        resultat = lire(mode="pages", debut_page=2, fin_page=4)
        assert "--- Page 2 ---" in resultat and "contenu4" in resultat
        assert "contenu5" not in resultat and "contenu1\n" not in resultat
        assert "Fin de l'intervalle demandé" in resultat

    def test_lecture_par_tranches_de_pages(self, stockage, monkeypatch):
        monkeypatch.setenv("DOCUMENT_PAGES_MAX_PAR_APPEL", "3")
        stockage["deposer"]("c.pdf", _pdf([f"contenu{i}" for i in range(1, 9)]), "application/pdf")
        resultat = lire(mode="texte")
        assert "contenu3" in resultat and "contenu4" not in resultat
        assert "debut_page=4" in resultat
        suite = lire(mode="texte", debut_page=7)
        assert "contenu8" in suite and "Fin du fichier" in suite

    def test_budget_de_caracteres(self, stockage, monkeypatch):
        monkeypatch.setenv("DOCUMENT_TRANCHE_CARACTERES", "120")
        stockage["deposer"]("c.pdf", _pdf(["x" * 60, "y" * 60, "z" * 60]), "application/pdf")
        resultat = lire(mode="texte")
        assert "debut_page=" in resultat and "zzzz" not in resultat

    def test_page_hors_du_fichier(self, stockage):
        stockage["deposer"]("c.pdf", _pdf(["a", "b"]), "application/pdf")
        assert "n'a que 2 pages" in lire(mode="pages", debut_page=9)

    def test_ocr_seulement_pour_les_pages_demandees(self, stockage, monkeypatch):
        stockage["deposer"]("scan.pdf", _pdf(["texte natif", "", "", "autre"]), "application/pdf")
        appels = []

        def faux_ocr(chemin, numero):
            appels.append(numero)
            return f"texte reconnu page {numero + 1}"

        monkeypatch.setattr("core.ocr_pages_scannees.extraire_texte_page_scannee", faux_ocr)
        resultat = lire(mode="pages", debut_page=2, fin_page=2)
        assert "texte reconnu page 2" in resultat
        assert appels == [1]

    def test_page_sans_texte_et_sans_ocr(self, stockage, monkeypatch):
        stockage["deposer"]("scan.pdf", _pdf(["", "ok"]), "application/pdf")
        monkeypatch.setattr("core.ocr_pages_scannees.extraire_texte_page_scannee", lambda c, n: None)
        assert "aucun texte lisible sur cette page" in lire(mode="pages", debut_page=1, fin_page=1)

    def test_une_page_illisible_ne_fait_pas_perdre_les_autres(self, stockage, monkeypatch):
        stockage["deposer"]("c.pdf", _pdf(["premiere", "", "derniere"]), "application/pdf")

        def ocr_qui_plante(chemin, numero):
            raise RuntimeError("tesseract absent")

        monkeypatch.setattr("core.ocr_pages_scannees.extraire_texte_page_scannee", ocr_qui_plante)
        resultat = lire(mode="texte")
        assert "premiere" in resultat and "derniere" in resultat

    def test_le_fichier_temporaire_est_supprime(self, stockage, monkeypatch):
        stockage["deposer"]("scan.pdf", _pdf([""]), "application/pdf")
        chemins = []

        def faux_ocr(chemin, numero):
            chemins.append(chemin)
            return "lu"

        monkeypatch.setattr("core.ocr_pages_scannees.extraire_texte_page_scannee", faux_ocr)
        lire(mode="texte")
        import os

        assert chemins and not os.path.exists(chemins[0])


class TestRecherche:
    def test_pdf_donne_la_page(self, stockage):
        stockage["deposer"]("c.pdf", _pdf(["rien", "la photosynthèse ici", "rien"]), "application/pdf")
        resultat = lire(mode="recherche", recherche="PHOTOSYNTHÈSE")
        assert "Page 2" in resultat and "Page 1" not in resultat

    def test_pages_scannees_signalees_et_non_lues(self, stockage, monkeypatch):
        stockage["deposer"]("c.pdf", _pdf(["mot", ""]), "application/pdf")
        ocr = mock.Mock()
        monkeypatch.setattr("core.ocr_pages_scannees.extraire_texte_page_scannee", ocr)
        resultat = lire(mode="recherche", recherche="mot")
        assert "1 pages sans texte natif" in resultat
        ocr.assert_not_called()

    def test_aucun_resultat(self, stockage):
        stockage["deposer"]("c.pdf", _pdf(["abc"]), "application/pdf")
        assert "Aucun passage" in lire(mode="recherche", recherche="zzz")

    def test_resultats_limites(self, stockage, monkeypatch):
        monkeypatch.setenv("DOCUMENT_RECHERCHE_RESULTATS_MAX", "2")
        stockage["deposer"]("t.txt", ("mot " * 10).encode(), "text/plain")
        resultat = lire(mode="recherche", recherche="mot")
        assert resultat.count("Position") == 2 and "limités à 2" in resultat

    def test_recherche_sans_mot(self, stockage):
        stockage["deposer"]("t.txt", b"abc", "text/plain")
        assert "demande le mot" in lire(mode="recherche")


class TestAutresFormats:
    def test_word(self, stockage):
        stockage["deposer"]("d.docx", _docx(["Premier paragraphe", "Deuxième paragraphe"]))
        resultat = lire(mode="texte")
        assert "Premier paragraphe" in resultat and "Fin du fichier" in resultat

    def test_texte_long_lu_par_tranches_sans_perte(self, stockage, monkeypatch):
        monkeypatch.setenv("DOCUMENT_TRANCHE_CARACTERES", "1000")
        texte = "".join(f"{i % 10}" for i in range(2500))
        stockage["deposer"]("long.txt", texte.encode(), "text/plain")
        premiere = lire(mode="texte")
        assert "debut_caractere=1000" in premiere
        deuxieme = lire(mode="texte", debut_caractere=1000)
        assert "debut_caractere=2000" in deuxieme
        derniere = lire(mode="texte", debut_caractere=2000)
        assert "Fin du fichier : 2500" in derniere

    def test_pas_de_limite_de_30000_caracteres_a_la_demande(self, stockage, monkeypatch):
        monkeypatch.setenv("DOCUMENT_TRANCHE_CARACTERES", "100000")
        stockage["deposer"]("long.txt", ("a" * 45000).encode(), "text/plain")
        resultat = lire(mode="texte")
        assert resultat.count("a") >= 45000 and "Fin du fichier" in resultat

    def test_excel(self, stockage):
        stockage["deposer"]("n.xlsx", _xlsx())
        resultat = lire(mode="texte")
        assert "Awa" in resultat and "Feuille : Moyennes" in resultat

    def test_pptx_par_diapositive(self, stockage):
        stockage["deposer"]("e.pptx", _pptx(["Titre un", "Titre deux", "Titre trois"]))
        resultat = lire(mode="pages", debut_page=2, fin_page=3)
        assert "Titre deux" in resultat and "Titre trois" in resultat and "Titre un" not in resultat
        assert "--- Diapositive 2 ---" in resultat

    def test_binaire_inconnu(self, stockage):
        stockage["deposer"]("appli.bin", bytes(range(256)) * 4, "application/octet-stream")
        assert "pas un format dont le texte peut être lu" in lire(mode="texte")

    def test_mode_pages_sur_un_format_sans_pages(self, stockage):
        stockage["deposer"]("t.txt", b"bonjour", "text/plain")
        assert "pas de pages" in lire(mode="pages")


class TestSecuriteEtErreurs:
    def test_fichier_introuvable_ou_autre_compte(self, stockage):
        stockage["deposer"]("c.pdf", _pdf(["a"]), "application/pdf")
        assert "introuvable" in lire(url_fichier="https://stockage/invente", mode="texte")

    def test_le_compte_de_l_utilisateur_est_transmis(self, stockage):
        stockage["deposer"]("c.pdf", _pdf(["a"]), "application/pdf")
        lire(mode="infos")
        assert stockage["appels"] == [("https://stockage/fichier", "u1")]

    def test_mode_inconnu(self, stockage):
        assert "Mode inconnu" in lire(mode="devine")

    def test_lien_manquant(self):
        assert "url_fichier manquant" in lecteur.lire_document("", "u1", "texte")

    def test_erreur_de_stockage_ne_leve_pas(self, monkeypatch):
        def plante(url, user_id):
            raise RuntimeError("R2 indisponible")

        monkeypatch.setattr(lecteur, "retrouver_fichier", plante)
        assert "réessaie" in lecteur.lire_document("https://x", "u1", "texte")

    def test_pdf_abime_ne_leve_pas(self, stockage):
        stockage["deposer"]("c.pdf", b"%PDF-1.4 abime", "application/pdf")
        assert "échoué" in lire(mode="texte")

    def test_aucun_double_tiret_dans_les_phrases_destinees_a_l_ia(self):
        for texte in lecteur.MESSAGES_RAISONS.values():
            assert " -- " not in texte


class TestDocumentJointPresent:
    BLOC = "\n\n[Document joint : cours.pdf]\nContenu non lu (PDF, 2,1 Mo, 40 pages).\n[Lien réel du fichier : https://x/y.pdf]"

    def test_nouveau_format_dans_le_message(self):
        assert lecteur.document_joint_present("regarde" + self.BLOC, [])

    def test_nouveau_format_dans_l_historique_utilisateur(self):
        assert lecteur.document_joint_present("et la suite ?", [{"role": "user", "content": "x" + self.BLOC}])

    def test_ancien_format_avec_texte_ne_declenche_pas(self):
        ancien = "\n\n[Document joint : cours.pdf]\nLe texte extrait du cours\n[Lien réel du fichier : https://x/y.pdf]"
        assert not lecteur.document_joint_present("regarde" + ancien, [])

    def test_message_de_l_assistant_ignore(self):
        assert not lecteur.document_joint_present("salut", [{"role": "assistant", "content": self.BLOC}])

    def test_sans_piece_jointe(self):
        assert not lecteur.document_joint_present("bonjour", None)


class TestLireFichierInchange:
    def test_limite_d_origine_conservee(self):
        resultat = extraction_documents.lire_fichier("t.txt", ("a" * 40000).encode(), "text/plain")
        assert resultat["tronque"] is True and len(resultat["texte"]) == extraction_documents.LONGUEUR_MAX_TEXTE_EXTRAIT

    def test_sans_limite(self):
        resultat = extraction_documents.lire_fichier("t.txt", ("a" * 40000).encode(), "text/plain", limite=None)
        assert resultat["tronque"] is False and len(resultat["texte"]) == 40000


class TestPasserelleVersLExecutionDeCode:
    @pytest.fixture
    def execution(self, monkeypatch):
        import core.outils_execution_code as code

        etat = {"ligne": {"nom_fichier": "Mon cours.pdf"}, "octets": b"%PDF contenu", "appels": []}

        def faux_retrouver(url, user_id):
            etat["appels"].append((url, user_id))
            if url != "https://stockage/fichier":
                return None
            return etat["ligne"], etat["octets"]

        monkeypatch.setattr("core.lecture_media_conversation.retrouver_fichier", faux_retrouver)
        etat["code"] = code
        return etat

    def test_destination_par_defaut(self, execution):
        prepare, erreur = execution["code"]._preparer_import("u1", "https://stockage/fichier", "")
        assert erreur is None and prepare == ("fichiers_joints/Mon cours.pdf", b"%PDF contenu")
        assert execution["appels"] == [("https://stockage/fichier", "u1")]

    def test_destination_choisie(self, execution):
        prepare, _ = execution["code"]._preparer_import("u1", "https://stockage/fichier", "travail/entree.pdf")
        assert prepare[0] == "travail/entree.pdf"

    def test_chemin_dangereux_refuse(self, execution):
        for chemin in ("../evasion.pdf", "/etc/passwd", "~/x"):
            prepare, erreur = execution["code"]._preparer_import("u1", "https://stockage/fichier", chemin)
            assert prepare is None and "Chemin refusé" in erreur

    def test_nom_d_origine_malveillant_neutralise(self, execution):
        execution["ligne"]["nom_fichier"] = "../../etc/passwd"
        prepare, _ = execution["code"]._preparer_import("u1", "https://stockage/fichier", "")
        assert prepare[0] == "fichiers_joints/passwd"

    def test_lien_inconnu_ou_autre_compte(self, execution):
        prepare, erreur = execution["code"]._preparer_import("u1", "https://stockage/invente", "")
        assert prepare is None and "introuvable" in erreur

    def test_lien_manquant(self, execution):
        prepare, erreur = execution["code"]._preparer_import("u1", "", "")
        assert prepare is None and "url_fichier manquant" in erreur

    def test_fichier_trop_lourd(self, execution, monkeypatch):
        monkeypatch.setenv("EXECUTION_CODE_IMPORT_TAILLE_MAX_MO", "1")
        execution["octets"] = b"x" * (2 * 1024 * 1024)
        prepare, erreur = execution["code"]._preparer_import("u1", "https://stockage/fichier", "")
        assert prepare is None and "au-delà de la limite" in erreur

    def test_erreur_de_stockage(self, execution, monkeypatch):
        def plante(url, user_id):
            raise RuntimeError("R2 indisponible")

        monkeypatch.setattr("core.lecture_media_conversation.retrouver_fichier", plante)
        prepare, erreur = execution["code"]._preparer_import("u1", "https://stockage/fichier", "")
        assert prepare is None and "réessaie" in erreur

    def test_ecriture_dans_le_workspace(self, execution):
        sandbox = mock.MagicMock()
        sandbox.commands.run.return_value = SimpleNamespace(exit_code=0, stdout="", stderr="")
        resultat = execution["code"]._executer_action(
            sandbox, "importer_fichier", "python", "", "", "", "", {}, ("fichiers_joints/cours.pdf", b"octets")
        )
        sandbox.files.write.assert_called_once_with("/home/user/fichiers_joints/cours.pdf", b"octets")
        assert "fichiers_joints/cours.pdf" in resultat and "original reste intact" in resultat

    def test_dossier_impossible_a_creer(self, execution):
        sandbox = mock.MagicMock()
        sandbox.commands.run.return_value = SimpleNamespace(exit_code=1, stdout="", stderr="")
        resultat = execution["code"]._executer_action(
            sandbox, "importer_fichier", "python", "", "", "", "", {}, ("a/b.pdf", b"x")
        )
        sandbox.files.write.assert_not_called()
        assert "Impossible de préparer" in resultat

    def test_sans_fichier_prepare(self, execution):
        resultat = execution["code"]._executer_action(mock.MagicMock(), "importer_fichier", "python", "", "", "", "", {})
        assert "Aucun fichier" in resultat

    def test_la_reservation_compte_une_ecriture_de_fichier(self, execution, monkeypatch):
        code = execution["code"]
        monkeypatch.setenv("E2B_API_KEY", "cle-de-test")
        monkeypatch.setenv("SUPABASE_SECRET", "secret")
        envoyes = []

        def faux_rpc(nom, parametres):
            envoyes.append((nom, parametres))
            return {"allowed": False, "reason": "daily_limit"}

        monkeypatch.setattr(code, "_rpc_objet", faux_rpc)
        monkeypatch.setattr(code, "_identite_interne", lambda ctx: ("u1", "c1", None))
        resultat = code.gerer_execution_code(
            action="importer_fichier", ctx=object(), url_fichier="https://stockage/fichier"
        )
        assert envoyes[0][0] == "execution_code_reserver"
        assert envoyes[0][1]["p_action"] == "write_file"
        assert "EXECUTION_CODE" in resultat

    def test_un_lien_invalide_ne_consomme_aucun_quota(self, execution, monkeypatch):
        code = execution["code"]
        monkeypatch.setenv("E2B_API_KEY", "cle-de-test")
        rpc = mock.Mock()
        monkeypatch.setattr(code, "_rpc_objet", rpc)
        monkeypatch.setattr(code, "_identite_interne", lambda ctx: ("u1", "c1", None))
        resultat = code.gerer_execution_code(action="importer_fichier", ctx=object(), url_fichier="https://stockage/invente")
        assert "introuvable" in resultat
        rpc.assert_not_called()
