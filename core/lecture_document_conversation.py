"""
Lecture à la demande d'un document joint à une conversation (27/10/2026,
demande Bourama : les fichiers joints marchent comme chez Claude, c'est l'IA
qui décide de les lire, à part les images qui restent lues à l'envoi).

Avant ce chantier, api/uploads.py:uploader_document_chat extrayait le texte de
chaque document DÈS l'envoi (OCR compris) et en injectait les 30000 premiers
caractères dans le message. Maintenant, l'envoi (route /api/uploads/fichier-chat)
enregistre seulement le fichier et ses informations peu coûteuses (type, taille,
nombre de pages). Ce module fait le vrai travail, quand l'IA appelle l'outil
lire_document_conversation (core/outils_lecture_document.py) :

- mode "infos" : informations seules, aucun coût
- mode "texte" : le texte, par tranches (par pages pour un PDF ou un PowerPoint,
  par caractères pour les autres formats)
- mode "pages" : un intervalle de pages ou de diapositives précis
- mode "recherche" : les passages qui contiennent un mot ou une expression

L'original n'est jamais modifié : on relit seulement ses octets. Le retour est
toujours un TEXTE destiné au modèle, jamais une exception : un document
illisible n'est pas une erreur de l'application, le modèle doit simplement le
dire à l'étudiant sans rien inventer.

Les limites (taille des tranches, pages par appel, résultats de recherche) viennent
de core/limites_media.py et se changent par variable d'environnement.
"""

import io
import logging
import os
import re
import tempfile

from core.extraction_documents import (
    _textes_formes_pptx,
    lire_fichier,
    type_document,
)
from core.lecture_media_conversation import formater_taille, retrouver_fichier
from core.limites_media import (
    pages_max_par_appel_document,
    resultats_max_recherche_document,
    tranche_document_caracteres,
)

MODES = ("infos", "texte", "pages", "recherche")

LIBELLES_FAMILLES = {
    "pdf": "PDF",
    "docx": "document Word",
    "xlsx": "tableur Excel",
    "pptx": "présentation PowerPoint",
    "office_ancien": "ancien format Office",
    "texte": "fichier texte ou code",
    "inconnu": "fichier",
}

# Largeur du contexte autour d'un mot trouvé par la recherche, de chaque côté.
CONTEXTE_RECHERCHE_CARACTERES = 90

MESSAGES_RAISONS = {
    "TYPE_NON_PRIS_EN_CHARGE": (
        "Ce fichier n'est pas un format dont le texte peut être lu (binaire). "
        "Dis-le à l'étudiant. Si un traitement par code est utile, copie-le dans l'espace "
        "d'exécution de code avec l'action importer_fichier de gerer_execution_code."
    ),
    "CONVERSION_INDISPONIBLE": (
        "Ce fichier est dans un ancien format Office et sa conversion n'est pas disponible "
        "pour le moment. Dis-le à l'étudiant, sans rien inventer sur son contenu."
    ),
    "AUCUN_TEXTE_TROUVE": (
        "Aucun texte n'a pu être lu dans ce fichier (vide, ou image sans texte reconnu). "
        "Dis-le à l'étudiant, sans rien inventer sur son contenu."
    ),
    "ECHEC_EXTRACTION": (
        "La lecture de ce fichier a échoué (fichier abîmé ou protégé). "
        "Dis-le à l'étudiant, sans rien inventer sur son contenu."
    ),
}


# Informations peu coûteuses, utilisées à l'envoi et par le mode infos

def infos_document(nom_fichier: str, contenu: bytes, type_mime: str = "") -> dict:
    """
    Informations d'un fichier joint sans lire son texte : famille, taille, et
    selon le format le nombre de pages (PDF), de diapositives (PowerPoint) ou
    les noms des feuilles (Excel). Ne lève jamais : une information qui ne peut
    pas être lue reste simplement absente (None).
    """
    famille = type_document(nom_fichier, type_mime)
    infos = {
        "famille": famille,
        "taille_octets": len(contenu),
        "pages": None,
        "unite": None,
        "feuilles": None,
    }
    try:
        if famille == "pdf":
            import PyPDF2

            infos["pages"] = len(PyPDF2.PdfReader(io.BytesIO(contenu)).pages)
            infos["unite"] = "pages"
        elif famille == "pptx":
            from pptx import Presentation

            infos["pages"] = len(Presentation(io.BytesIO(contenu)).slides)
            infos["unite"] = "diapositives"
        elif famille == "xlsx":
            import openpyxl

            classeur = openpyxl.load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
            infos["feuilles"] = list(classeur.sheetnames)
            classeur.close()
    except Exception as e:
        logging.warning(f"Informations de document illisibles ({nom_fichier}) : {e}")
    return infos


def resume_infos(infos: dict) -> str:
    """Résumé court des informations, le même texte pour le message et l'outil."""
    morceaux = [LIBELLES_FAMILLES.get(infos.get("famille"), "fichier"), formater_taille(infos.get("taille_octets") or 0)]
    if infos.get("pages"):
        morceaux.append(f"{infos['pages']} {infos.get('unite') or 'pages'}")
    if infos.get("feuilles"):
        feuilles = infos["feuilles"]
        morceaux.append(f"{len(feuilles)} feuille{'s' if len(feuilles) > 1 else ''} : " + ", ".join(feuilles[:8]))
    return ", ".join(morceaux)


# Lecture page par page (PDF et PowerPoint)

class _LecteurPdf:
    """
    Lit un PDF page par page. Une page sans texte natif (scan, photo) passe par
    le même OCR de secours que l'ancienne lecture à l'envoi (Tesseract d'abord,
    Gemini vision en filet de sécurité, voir core/ocr_pages_scannees.py), mais
    SEULEMENT pour les pages réellement demandées.
    """

    def __init__(self, contenu: bytes):
        import PyPDF2

        self._contenu = contenu
        self._lecteur = PyPDF2.PdfReader(io.BytesIO(contenu))
        self.nombre_pages = len(self._lecteur.pages)
        self._chemin_temp = None

    def texte_natif(self, index: int) -> str:
        return (self._lecteur.pages[index].extract_text() or "").strip()

    def texte(self, index: int) -> str:
        texte = self.texte_natif(index)
        if texte:
            return texte
        from core.ocr_pages_scannees import extraire_texte_page_scannee

        if self._chemin_temp is None:
            with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
                tmp.write(self._contenu)
                self._chemin_temp = tmp.name
        return (extraire_texte_page_scannee(self._chemin_temp, index) or "").strip()

    def fermer(self):
        if self._chemin_temp:
            try:
                os.remove(self._chemin_temp)
            except OSError:
                pass
            self._chemin_temp = None


class _LecteurPptx:
    """Lit une présentation diapositive par diapositive (texte des formes et notes)."""

    def __init__(self, contenu: bytes):
        from pptx import Presentation

        self._diapositives = list(Presentation(io.BytesIO(contenu)).slides)
        self.nombre_pages = len(self._diapositives)

    def texte_natif(self, index: int) -> str:
        return self.texte(index)

    def texte(self, index: int) -> str:
        diapositive = self._diapositives[index]
        morceaux = list(_textes_formes_pptx(diapositive.shapes))
        if diapositive.has_notes_slide:
            notes = diapositive.notes_slide.notes_text_frame.text
            if notes.strip():
                morceaux.append(f"Notes : {notes}")
        return "\n".join(morceaux).strip()

    def fermer(self):
        return None


def _lire_pages(lecteur, debut_page: int, fin_page: int, unite: str) -> str:
    """Texte d'un intervalle de pages, borné en nombre de pages ET en caractères."""
    total = lecteur.nombre_pages
    debut = max(1, int(debut_page or 1))
    fin = int(fin_page or 0)
    fin = total if fin <= 0 else min(fin, total)
    if debut > total:
        return f"Ce fichier n'a que {total} {unite} : {unite[:-1] if unite.endswith('s') else unite} {debut} n'existe pas."
    if fin < debut:
        return f"L'intervalle est vide : la fin ({fin}) est avant le début ({debut})."

    pages_max = pages_max_par_appel_document()
    budget = tranche_document_caracteres()
    blocs = []
    utilises = 0
    derniere_lue = debut - 1
    for numero in range(debut, fin + 1):
        if numero - debut >= pages_max:
            break
        try:
            texte = lecteur.texte(numero - 1)
        except Exception as e:
            # Une page illisible ne doit pas faire perdre les pages déjà lues.
            logging.warning(f"Page {numero} illisible : {e}")
            texte = ""
        etiquette = f"--- {unite[:-1].capitalize() if unite.endswith('s') else unite.capitalize()} {numero} ---"
        if not texte:
            bloc = f"{etiquette}\n(aucun texte lisible sur cette page)"
        else:
            bloc = f"{etiquette}\n{texte}"
        if blocs and utilises + len(bloc) > budget:
            break
        if not blocs and len(bloc) > budget:
            bloc = bloc[:budget] + "\n[page tronquée : trop de texte pour un seul appel]"
        blocs.append(bloc)
        utilises += len(bloc)
        derniere_lue = numero

    resultat = "\n\n".join(blocs)
    if derniere_lue < fin:
        resultat += (
            f"\n\n[Lecture partielle : {unite} {debut} à {derniere_lue} sur {total}. "
            f"Rappelle l'outil avec debut_page={derniere_lue + 1} pour la suite, "
            "avant de conclure si tu as besoin du reste.]"
        )
    elif derniere_lue < total:
        resultat += f"\n\n[Fin de l'intervalle demandé. Le fichier a {total} {unite}.]"
    else:
        resultat += f"\n\n[Fin du fichier ({total} {unite}).]"
    return resultat


def _extraits(texte: str, mot: str, limite: int):
    """Les passages d'un texte qui contiennent `mot` (sans tenir compte de la casse)."""
    bas = texte.casefold()
    cherche = mot.casefold()
    resultats = []
    depart = 0
    while len(resultats) < limite:
        position = bas.find(cherche, depart)
        if position < 0:
            break
        debut = max(0, position - CONTEXTE_RECHERCHE_CARACTERES)
        fin = min(len(texte), position + len(mot) + CONTEXTE_RECHERCHE_CARACTERES)
        extrait = " ".join(texte[debut:fin].split())
        resultats.append((position, ("..." if debut > 0 else "") + extrait + ("..." if fin < len(texte) else "")))
        depart = position + max(1, len(mot))
    return resultats


def _rechercher_pages(lecteur, mot: str, unite: str) -> str:
    """Recherche dans les pages sans OCR : seules les pages à texte natif sont fouillées."""
    maximum = resultats_max_recherche_document()
    lignes = []
    sans_texte = 0
    trouves = 0
    singulier = unite[:-1] if unite.endswith("s") else unite
    for index in range(lecteur.nombre_pages):
        if trouves >= maximum:
            break
        try:
            texte = lecteur.texte_natif(index)
        except Exception:
            texte = ""
        if not texte:
            sans_texte += 1
            continue
        for _, extrait in _extraits(texte, mot, maximum - trouves):
            lignes.append(f"{singulier.capitalize()} {index + 1} : {extrait}")
            trouves += 1
            if trouves >= maximum:
                break
    if not lignes:
        resultat = f"Aucun passage ne contient « {mot} »."
    else:
        resultat = "\n".join(lignes)
        if trouves >= maximum:
            resultat += f"\n\n[Résultats limités à {maximum}. Précise la recherche ou lis les {unite} concernées.]"
    if sans_texte:
        resultat += (
            f"\n\n[{sans_texte} {unite} sans texte natif (scan ou image) n'ont pas été fouillées. "
            f"Lis-les avec le mode pages si tu en as besoin.]"
        )
    return resultat


# Lecture par caractères (Word, Excel, texte, code, anciens formats)

def _lire_caracteres(texte: str, debut_caractere: int) -> str:
    total = len(texte)
    debut = max(0, int(debut_caractere or 0))
    if debut >= total:
        return f"Le texte ne contient que {total} caractères : la position {debut} est après la fin."
    tranche = tranche_document_caracteres()
    fin = min(total, debut + tranche)
    resultat = texte[debut:fin]
    if fin < total:
        resultat += (
            f"\n\n[Lecture partielle : caractères {debut} à {fin} sur {total}. "
            f"Rappelle l'outil avec debut_caractere={fin} pour la suite, "
            "avant de conclure si tu as besoin du reste.]"
        )
    else:
        resultat += f"\n\n[Fin du fichier : {total} caractères au total.]"
    return resultat


def _rechercher_caracteres(texte: str, mot: str) -> str:
    maximum = resultats_max_recherche_document()
    resultats = _extraits(texte, mot, maximum)
    if not resultats:
        return f"Aucun passage ne contient « {mot} »."
    lignes = [f"Position {position} : {extrait}" for position, extrait in resultats]
    resultat = "\n".join(lignes)
    if len(resultats) >= maximum:
        resultat += f"\n\n[Résultats limités à {maximum}. Précise la recherche ou lis le passage avec debut_caractere.]"
    return resultat


# Bloc écrit par le frontend (lib/preparerPiecesJointes.ts) pour un document dont le
# contenu n'a PAS été lu à l'envoi. Les anciens messages, qui portent le texte du
# document à la place de cette mention, ne le déclenchent pas.
MOTIF_DOCUMENT_NON_LU = re.compile(r"\[Document joint : [^\]\n]+\]\nContenu non lu")


def document_joint_present(message_utilisateur, historique) -> bool:
    """
    True si le message courant ou un message de l'UTILISATEUR de l'historique
    contient un document joint non lu. Sert à main.py pour donner l'outil
    lire_document_conversation au modèle : sans cela, rien ne garantit qu'il soit
    proposé ce tour-ci, alors que le message lui dit que le contenu n'est pas lu.
    """
    textes = [message_utilisateur or ""]
    for m in historique or []:
        if isinstance(m, dict) and m.get("role") == "user":
            textes.append(m.get("content") or "")
    return any(MOTIF_DOCUMENT_NON_LU.search(texte) for texte in textes)


def lire_document(
    url_fichier: str,
    user_id: str,
    mode: str = "infos",
    debut_page: int = 0,
    fin_page: int = 0,
    debut_caractere: int = 0,
    recherche: str = "",
) -> str:
    """
    Lit un document joint à la conversation de cet utilisateur, retrouvé par son
    lien réel. Voir le haut de ce fichier pour les modes. Renvoie toujours un texte.
    """
    mode = (mode or "infos").strip().lower()
    if mode not in MODES:
        return f"Mode inconnu : « {mode} ». Modes possibles : infos, texte, pages, recherche."
    url_fichier = (url_fichier or "").strip()
    if not url_fichier:
        return "url_fichier manquant (le lien réel indiqué juste après le fichier joint)."
    if mode == "recherche" and not (recherche or "").strip():
        return "Le mode recherche demande le mot ou l'expression à chercher dans `recherche`."

    try:
        trouve = retrouver_fichier(url_fichier, user_id)
    except Exception as e:
        logging.error(f"ERREUR relecture du document joint : {e}")
        return "Impossible de relire ce fichier pour le moment, réessaie."
    if trouve is None:
        return "Ce fichier est introuvable, ou ne t'appartient pas."
    ligne, contenu = trouve
    nom = ligne.get("nom_fichier") or "document"
    type_mime = ligne.get("type_mime") or ""

    infos = infos_document(nom, contenu, type_mime)
    entete = f"Document : {nom} ({resume_infos(infos)})"
    if mode == "infos":
        return entete

    famille = infos["famille"]
    lecteur = None
    try:
        if famille in ("pdf", "pptx"):
            try:
                lecteur = _LecteurPdf(contenu) if famille == "pdf" else _LecteurPptx(contenu)
            except Exception as e:
                logging.error(f"ERREUR ouverture du document joint ({nom}) : {e}")
                return entete + "\n" + MESSAGES_RAISONS["ECHEC_EXTRACTION"]
            unite = "pages" if famille == "pdf" else "diapositives"
            if mode == "recherche":
                corps = _rechercher_pages(lecteur, recherche.strip(), unite)
            else:
                corps = _lire_pages(lecteur, debut_page, fin_page, unite)
            return entete + "\n\n" + corps

        # Les autres formats n'ont pas de pages : lecture du texte complet (sans la
        # limite d'origine de 30000 caractères), puis par tranches.
        lecture = lire_fichier(nom, contenu, type_mime, limite=None)
        texte = lecture["texte"]
        if texte is None:
            return entete + "\n" + MESSAGES_RAISONS.get(lecture["raison"], MESSAGES_RAISONS["ECHEC_EXTRACTION"])
        if mode == "recherche":
            return entete + "\n\n" + _rechercher_caracteres(texte, recherche.strip())
        note = ""
        if mode == "pages":
            note = "(Ce format n'a pas de pages : lecture par caractères.)\n\n"
        return entete + "\n\n" + note + _lire_caracteres(texte, debut_caractere)
    except Exception as e:
        logging.error(f"ERREUR lecture du document joint ({nom}) : {e}")
        return entete + "\n" + MESSAGES_RAISONS["ECHEC_EXTRACTION"]
    finally:
        if lecteur is not None:
            lecteur.fermer()
