"""
Module central d'extraction de contenu (bibliothèque privée, bibliothèque
publique, dossiers téléphone).

Chantier du 18/09/2026 (demande Bourama, suite au signalement "plein
d'éléments que Clovis prétend ne pas pouvoir lire") : jusqu'ici, chaque
bibliothèque réimplémentait sa propre dispatch "quel extracteur pour quel
type de fichier", avec une couverture différente et divergente selon
l'endroit -- notamment Word/Excel/vidéo jamais vectorisés côté PRIVÉ
(core/file_attente_vectorisation.py::_vectoriser_privee, avant ce
chantier), alors que le catalogue public et les dossiers désignés
savaient déjà le faire à la demande. Bourama a explicitement demandé
d'arrêter de construire type par type, système par système : ce module
réunit UN SEUL point d'entrée, extraire_segments(), qui route vers le
bon extracteur déjà existant dans le dépôt (rien réécrit depuis zéro --
PDF/image/audio/vidéo réutilisent tels quels core.bibliotheque_rag et
core.description_multimedia). Ajouter un type de fichier plus tard =
une seule branche ajoutée ici, profite à toute bibliothèque qui appelle
ce module.

Ne fait QUE l'extraction (texte / description / segments horodatés ou
paginés) -- jamais l'indexation (embeddings, écriture des chunks en
base), qui reste propre à chaque bibliothèque (documents_bibliotheque /
documents_catalogue_public / dossiers désignés) pour ne rien casser de
leur système de citations cliquables respectif. Chaque segment renvoyé
porte, quand elle a un sens, sa position d'origine (page_debut/page_fin
pour un PDF, timestamp_debut/timestamp_fin pour un audio/vidéo) --
directement passable en **kwargs (sans "texte") à
core.bibliotheque_rag.indexer_texte_bibliotheque.

Étapes 1+2 de ce chantier (ce module + branchement bibliothèque privée)
seulement pour l'instant -- le catalogue public et les dossiers désignés
(téléphone) gardent leur système actuel, pas encore branchés ici (voir
CONTEXTE_extraction_contenu_bibliotheques.md pour la suite prévue).
"""

import os
import tempfile

from core.bibliotheque_rag import extraire_pages_pdf
from core.description_multimedia import (
    decrire_image_bibliotheque,
    transcrire_audio_bibliotheque,
    transcrire_et_decrire_video_bibliotheque,
)

TYPE_MIME_WORD = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TYPE_MIME_EXCEL = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Extensions traitées comme texte brut quand le type MIME est absent ou
# générique (ex. "application/octet-stream", fréquent pour un .md selon
# l'OS/le navigateur qui a fait l'upload) -- sans ça, ces fichiers ne
# seraient jamais reconnus comme extractibles du tout.
_EXTENSIONS_TEXTE_CONNUES = {
    "md", "markdown", "txt", "csv", "tsv", "json", "yml", "yaml", "xml",
    "py", "js", "jsx", "ts", "tsx", "html", "htm", "css", "scss", "sql",
    "sh", "bash", "java", "c", "cpp", "h", "hpp", "go", "rs", "rb", "php",
    "ini", "toml", "log",
}


def _extension(nom_fichier: str) -> str:
    nom_fichier = nom_fichier or ""
    return nom_fichier.rsplit(".", 1)[-1].lower() if "." in nom_fichier else ""


def extraire_texte_docx(contenu: bytes) -> str:
    """
    Texte + tableaux d'un .docx. Même logique que les copies déjà
    existantes dans le dépôt (api/uploads.py, core/file_attente_
    vectorisation.py, core/vectorisation_dossiers_designes.py) --
    celles-ci ne sont pas touchées par les étapes 1+2 de ce chantier,
    seule la bibliothèque privée délègue ici pour l'instant (voir
    fichier de contexte).
    """
    import io
    import docx

    document = docx.Document(io.BytesIO(contenu))
    morceaux = [p.text for p in document.paragraphs]
    for table in document.tables:
        for ligne in table.rows:
            morceaux.append("\t".join(cellule.text for cellule in ligne.cells))
    return "\n".join(morceaux)


def extraire_texte_xlsx(contenu: bytes) -> str:
    """Même principe que extraire_texte_docx, pour un .xlsx."""
    import io
    import openpyxl

    classeur = openpyxl.load_workbook(io.BytesIO(contenu), data_only=True)
    morceaux = []
    for feuille in classeur.worksheets:
        morceaux.append(f"--- Feuille : {feuille.title} ---")
        for ligne in feuille.iter_rows(values_only=True):
            morceaux.append("\t".join("" if v is None else str(v) for v in ligne))
    return "\n".join(morceaux)


def type_extractible(type_mime: str | None, nom_fichier: str = "") -> bool:
    """
    True si ce fichier a un extracteur connu ici -- donc si on peut
    raisonnablement s'attendre à en tirer au moins un chunk de texte.

    Sert à décider si un type doit être mis "en_attente" de
    vectorisation automatique (voir necessite_vectorisation_fichier_
    privee dans core/file_attente_vectorisation.py) : un type NON
    extractible ici (zip, exe, binaire inconnu...) ne doit JAMAIS être
    mis en attente -- ça finirait en échec de vectorisation trompeur (0
    chunk produit, voir _a_produit_des_chunks) plutôt que de rester
    simplement "pret" (rien à en attendre, comportement normal pour ce
    type).
    """
    type_mime = (type_mime or "").strip().lower()
    if type_mime == "application/pdf":
        return True
    if type_mime.startswith(("image/", "audio/", "video/", "text/")):
        return True
    if type_mime in (TYPE_MIME_WORD, TYPE_MIME_EXCEL, "application/json"):
        return True
    # Type MIME absent ou générique (souvent le cas pour un .md/.py/...
    # selon l'OS/le navigateur qui a fait l'upload) : se rabat sur
    # l'extension si elle est du texte reconnu.
    if type_mime in ("", "application/octet-stream"):
        return _extension(nom_fichier) in _EXTENSIONS_TEXTE_CONNUES
    return False


def _extraire_segments_pdf(contenu: bytes) -> list[dict]:
    chemin_temp = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
            tmp.write(contenu)
            chemin_temp = tmp.name
        segments = []
        for numero, texte_page in enumerate(extraire_pages_pdf(chemin_temp), start=1):
            if texte_page.strip():
                segments.append({"texte": texte_page, "page_debut": numero, "page_fin": numero})
        return segments
    finally:
        if chemin_temp:
            try:
                os.remove(chemin_temp)
            except OSError:
                pass


def _segments_depuis_transcription(segments_audio: list[dict] | None) -> list[dict]:
    resultat = []
    for segment in segments_audio or []:
        texte = (segment.get("text") or "").strip()
        if texte:
            resultat.append({
                "texte": texte,
                "timestamp_debut": segment.get("start"),
                "timestamp_fin": segment.get("end"),
            })
    return resultat


def extraire_segments(contenu: bytes, type_mime: str | None, nom_fichier: str = "") -> list[dict]:
    """
    Point d'entrée unique : extrait le contenu exploitable d'un fichier,
    quel que soit son type, sous forme de segments prêts à indexer.

    Chaque segment : {"texte": str, "page_debut": int|None,
    "page_fin": int|None, "timestamp_debut": float|None,
    "timestamp_fin": float|None} (les clés de position absentes valent
    None via .get() côté appelant -- pas la peine de toutes les mettre
    à chaque fois ici).

    Renvoie [] si le type n'est pas reconnu ou si rien n'a pu être
    extrait (fichier vide, description/transcription en échec...) --
    à l'appelant de décider quoi faire d'une liste vide.
    """
    type_mime_norm = (type_mime or "").strip().lower()

    if type_mime_norm == "application/pdf":
        return _extraire_segments_pdf(contenu)

    if type_mime_norm.startswith("image/"):
        description = decrire_image_bibliotheque(contenu, type_mime_norm)
        return [{"texte": description}] if description and description.strip() else []

    if type_mime_norm.startswith("audio/"):
        return _segments_depuis_transcription(transcrire_audio_bibliotheque(contenu, nom_fichier))

    if type_mime_norm.startswith("video/"):
        extension = _extension(nom_fichier) or "mp4"
        resultat = transcrire_et_decrire_video_bibliotheque(contenu, nom_fichier, extension) or {}
        segments = _segments_depuis_transcription(resultat.get("segments_audio"))
        for description in resultat.get("descriptions_frames") or []:
            if description and description.strip():
                segments.append({"texte": description})
        return segments

    if type_mime_norm == TYPE_MIME_WORD:
        texte = extraire_texte_docx(contenu)
        return [{"texte": texte}] if texte.strip() else []

    if type_mime_norm == TYPE_MIME_EXCEL:
        texte = extraire_texte_xlsx(contenu)
        return [{"texte": texte}] if texte.strip() else []

    if type_mime_norm.startswith("text/") or type_mime_norm == "application/json" or (
        type_mime_norm in ("", "application/octet-stream") and _extension(nom_fichier) in _EXTENSIONS_TEXTE_CONNUES
    ):
        texte = contenu.decode("utf-8", errors="ignore")
        return [{"texte": texte}] if texte.strip() else []

    return []
